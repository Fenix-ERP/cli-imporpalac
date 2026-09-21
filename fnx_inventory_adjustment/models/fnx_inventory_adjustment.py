import base64
import csv
import io

import openpyxl
import xlsxwriter

from odoo import Command, _, api, fields, models
from odoo.exceptions import UserError


class FnxInventoryAdjustment(models.Model):
    _name = "fnx.inventory.adjustment"
    _description = "FnxInventoryAdjustment"

    name = fields.Char(
        string="Reference",
        required=True,
        copy=False,
        readonly=True,
        default="/",
    )

    company_id = fields.Many2one(
        "res.company",
        default=lambda self: self.env.company.id,
    )

    date = fields.Date(
        default=lambda self: fields.Date.context_today(self),
        required=True,
    )

    excel_file = fields.Binary(
        required=True,
    )

    excel_filename = fields.Char()

    state = fields.Selection(
        [
            ("draft", "Draft"),
            ("confirmed", "Confirmed"),
            ("done", "Done"),
            ("cancel", "Cancel"),
        ],
        default="draft",
    )

    adjustment_type = fields.Selection(
        [
            ("cost", "Adjust Global Cost"),
            ("qty", "Adjust Quantity"),
            ("cost_qty", "Adjust Quantity and Cost"),
        ],
        default="cost",
        required=True,
    )

    reason_id = fields.Many2one(
        "fnx.adjustment.reason",
        required=True,
    )

    origin_location_id = fields.Many2one(
        "stock.location",
        required=True,
        domain=[
            (
                "usage",
                "in",
                [
                    "internal",
                ],
            )
        ],
    )

    notes = fields.Text()
    terms_and_conditions = fields.Text()

    user_id = fields.Many2one(
        "res.users",
        default=lambda self: self.env.user.id,
        string="Responsible",
        required=True,
    )

    line_ids = fields.One2many(
        "fnx.inventory.adjustment.line",
        "adjustment_id",
    )

    # account_id of "stock.valuation.layer.revaluation"
    account_id = fields.Many2one(
        "account.account",
    )

    # pickings for cost_qty adjustment_type
    picking_ids = fields.Many2many("stock.picking")

    account_move_ids = fields.Many2many(
        "account.move",
        string="Accounting Entries",
        copy=False,
    )

    def action_open_pickings(self):
        self.ensure_one()
        action = {
            "name": _("Adjustment Pickings"),
            "type": "ir.actions.act_window",
            "res_model": "stock.picking",
            "view_mode": "tree,form",
            "domain": [("id", "in", self.picking_ids.ids)],
        }
        return action

    @api.model_create_multi
    def create(self, vals_list):

        return super().create(vals_list)

    def export_template(self):
        option_map = {
            "qty": {
                "value": "qty",
                "label": _("Quantity"),
            },
            "cost": {
                "value": "cost",
                "label": _("Cost"),
            },
        }
        for record in self:
            try:
                output = io.BytesIO()
                workbook = xlsxwriter.Workbook(
                    output,
                    {
                        "in_memory": True,
                    },
                )
                worksheet = workbook.add_worksheet("Template")

                header_format = workbook.add_format(
                    {
                        "bold": True,
                        "text_wrap": True,
                        "valign": "vcenter",
                        "align": "center",
                        "border": 1,
                        "bg_color": "#D9E1F2",
                    }
                )
                text_format = workbook.add_format(
                    {
                        "valign": "vcenter",
                        "align": "left",
                        "border": 1,
                    }
                )
                number_format = workbook.add_format(
                    {
                        "valign": "vcenter",
                        "align": "right",
                        "border": 1,
                        "num_format": "#,##0.00",
                    }
                )
                header = [
                    _("Product Code"),
                ]
                if record.line_ids:
                    header.append(
                        _("System Qty (Dont Fill)"),
                    )
                header.append(
                    _("New Qty"),
                )
                if record.line_ids:
                    header.append(
                        _("ACtual Cost (Dont Fill)"),
                    )

                header.append(
                    _("New Cost"),
                )
                header.append(
                    _("Adjustment Type"),
                )

                technical_fields = (
                    [
                        "product_code",
                        "system_qty",
                        "new_qty",
                        "system_cost",
                        "real_cost",
                    ]
                    if record.line_ids
                    else ["product_code", "new_qty", "real_cost"]
                )
                for col_num, tech_field in enumerate(technical_fields):
                    worksheet.write(1, col_num, tech_field, text_format)

                for col_num, header in enumerate(header):
                    worksheet.write(0, col_num, header, header_format)

                worksheet.set_column(0, 5, 25)

                options_list = [str(opt["label"]) for opt in option_map.values()]
                max_row = max(1000, len(record.line_ids) + 100)
                col_validation = 5 if record.line_ids else 3
                worksheet.data_validation(
                    2,
                    col_validation,
                    max_row,
                    col_validation,
                    {
                        "validate": "list",
                        "source": options_list,
                        "input_title": _("Adjustment Type"),
                        "input_message": _("Select an option"),
                        "error_title": _("No valid value"),
                        "error_message": _("Select an option in the list %s")
                        % ", ".join(options_list),
                    },
                )

                row = 2
                for line in record.line_ids:
                    code = line.product_id.default_code or line.product_id.name or ""
                    worksheet.write(row, 0, code, text_format)
                    (
                        worksheet.write(
                            row, 1, line.product_id.qty_available, text_format
                        )
                        if record.line_ids
                        else None
                    )
                    worksheet.write(
                        row, 2 if record.line_ids else 1, line.new_qty, number_format
                    )
                    (
                        worksheet.write(
                            row, 3, line.product_id.standard_price, number_format
                        )
                        if record.line_ids
                        else None
                    )
                    worksheet.write(
                        row, 4 if record.line_ids else 3, line.real_cost, number_format
                    )
                    row += 1

                workbook.close()
                file_data = base64.b64encode(output.getvalue())
                output.close()

                attachment = self.env["ir.attachment"].create(
                    {
                        "name": "plantilla_ajuste_inventario.xlsx",
                        "type": "binary",
                        "datas": file_data,
                        "res_model": record._name,
                        "res_id": record.id if isinstance(record.id, int) else 0,
                        "mimetype": (
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                        ),
                    }
                )

                return {
                    "type": "ir.actions.act_url",
                    "url": f"/web/content/{attachment.id}?download=true",
                    "target": "self",
                }
            except Exception as e:
                record.env.user.notify_danger(_("Error: %s") % str(e))

    @staticmethod
    def _safe_float(val):
        try:
            return float(val or 0.0)
        except (ValueError, TypeError):
            return 0.0

    def _read_excel_file(self, excel_file):
        file_data = base64.b64decode(excel_file)
        wb = openpyxl.load_workbook(io.BytesIO(file_data), data_only=True)
        ws = wb.active

        fieldnames = [cell.value for cell in ws[2]]

        output = io.StringIO()
        writer = csv.writer(output)
        for row in ws.iter_rows(min_row=3, values_only=True):
            writer.writerow(row)

        output.seek(0)
        reader = csv.DictReader(output, fieldnames=fieldnames)

        result = []
        for row_dict in reader:
            code = row_dict.get("product_code", False)
            if not code:
                continue
            result.append(
                {
                    "code": code,
                    "new_qty": self._safe_float(row_dict.get("new_qty", 0)),
                    "real_cost": self._safe_float(row_dict.get("real_cost", 0)),
                }
            )
        return result

    def _prepare_line_commands(self, excel_data_list):
        line_commands = []
        excel_dict = {row["code"]: row for row in excel_data_list}

        for line in self.line_ids:
            code = line.product_id.default_code
            if code in excel_dict:
                row = excel_dict.pop(code)
                line_commands.append(
                    (
                        1,
                        line.id,
                        {
                            "new_qty": row.get("new_qty", 0),
                            "real_cost": row.get("real_cost", 0),
                        },
                    )
                )

        if excel_dict:
            products = self.env["product.product"].search(
                [("default_code", "in", list(excel_dict.keys()))]
            )
            for prod in products:
                row = excel_dict.pop(prod.default_code)
                line_commands.append(
                    (
                        0,
                        0,
                        {
                            "product_id": prod.id,
                            "new_qty": row.get("new_qty", 0),
                            "real_cost": row.get("real_cost", 0),
                        },
                    )
                )

        return line_commands

    def import_template(self):
        for record in self:
            if not record.excel_file:
                record.env.user.notify_warning(
                    _("Please upload an Excel file first in the 'Excel File' field."),
                    title=_("File required"),
                )
                continue

            try:
                excel_data = record._read_excel_file(record.excel_file)
            except Exception as e:
                record.env.user.notify_danger(
                    _("Error reading Excel file: %s") % str(e),
                    title=_("Read Error"),
                )
                continue

            line_commands = record._prepare_line_commands(excel_data)
            if line_commands:
                record.write({"line_ids": line_commands})
                record.env.user.notify_success(
                    _("Lines imported successfully."),
                    title=_("Success"),
                )
            else:
                record.env.user.notify_warning(
                    _("No matching products found to update or create."),
                    title=_("No Data"),
                )

    def action_confirm(self):
        code_map = {
            "cost": "fnx.cost.inventory.adjustment.sequence",
            "qty": "fnx.qty.inventory.adjustment.sequence",
            "cost_qty": "fnx.qty.cost.inventory.adjustment.sequence",
        }
        for record in self:
            if record.name == "/":
                record.name = (
                    self.env["ir.sequence"].next_by_code(
                        code_map.get(record.adjustment_type),
                    )
                    or "/"
                )

            record.write(dict(state="confirmed"))

    def action_cancel(self):
        for record in self:
            record.write(
                dict(
                    state="cancel",
                ),
            )

    def action_process(self):
        for record in self:
            lines = record.line_ids

            # Qty Process
            if record.adjustment_type == "qty":
                lines._process_qty()

            # Cost Process
            if record.adjustment_type == "cost":
                lines._process_cost()

            # Qty and Cost Process
            if record.adjustment_type == "cost_qty":
                lines._process_qty_cost()

            record.write(
                dict(
                    state="done",
                ),
            )


class FnxInventoryAdjustmentLine(models.Model):
    _name = "fnx.inventory.adjustment.line"
    _description = "FnxInventoryAdjustmentLine"

    stock_quant_id = fields.Many2one(
        "stock.quant",
    )

    adjustment_id = fields.Many2one(
        "fnx.inventory.adjustment",
    )

    stock_move_ids = fields.Many2many(
        "stock.move",
        string="Stock Moves",
        copy=False,
    )

    account_move_ids = fields.Many2many(
        "account.move",
        string="Accounting Entries",
        copy=False,
    )

    product_id = fields.Many2one(
        "product.product",
        required=True,
    )

    currency_id = fields.Many2one(
        "res.currency",
        related="product_id.currency_id",
    )

    system_qty = fields.Float(
        compute="_compute_system_qty",
    )

    new_qty = fields.Float()

    qty_var = fields.Float(
        compute="_compute_qty_var",
    )

    product_uom_id = fields.Many2one(
        "uom.uom",
        related="product_id.uom_id",
    )

    system_cost = fields.Float(
        related="product_id.standard_price",
    )

    real_cost = fields.Float()

    parent_state = fields.Selection(
        related="adjustment_id.state",
    )

    @api.depends("new_qty")
    def _compute_qty_var(self):
        for record in self:
            record.qty_var = record.new_qty - record.system_qty

    @api.depends("product_id", "adjustment_id.origin_location_id")
    def _compute_system_qty(self):
        for record in self:
            quant = self.env["stock.quant"].search(
                [
                    ("product_id", "=", record.product_id.id),
                    ("location_id", "=", record.adjustment_id.origin_location_id.id),
                ],
            )
            record.system_qty = sum(quant.mapped("quantity"))

    def _process_qty(self):
        for line in self:
            inventory_name = _(
                "Inventory Adjustment %(adjustment)s - %(line)s",
                adjustment=line.adjustment_id.name,
                line=line.id,
            )
            StockQuant = self.env["stock.quant"].with_context(
                inventory_mode=True,
                inventory_name=inventory_name,
            )

            quant = StockQuant.search(
                [
                    ("product_id", "=", line.product_id.id),
                    ("location_id", "=", line.adjustment_id.origin_location_id.id),
                ],
                limit=1,
            )

            if quant:
                quant.inventory_quantity = line.new_qty
            else:
                quant = StockQuant.create(
                    {
                        "product_id": line.product_id.id,
                        "location_id": line.adjustment_id.origin_location_id,
                        "inventory_quantity": line.new_qty,
                    }
                )

            quant.action_apply_inventory()

            moves = self.env["stock.move"].search(
                [
                    ("name", "=", inventory_name),
                    ("product_id", "=", line.product_id.id),
                    ("is_inventory", "=", True),
                ],
                order="id desc",
                limit=1,
            )
            account_moves = moves.mapped("account_move_ids") | moves.mapped(
                "stock_valuation_layer_ids.account_move_id"
            )
            if moves:
                line.stock_move_ids = [Command.link(move.id) for move in moves]
            if account_moves:
                line.account_move_ids = [
                    Command.link(move.id) for move in account_moves
                ]
                line.adjustment_id.account_move_ids = [
                    Command.link(move.id) for move in account_moves
                ]

    def _process_cost(self):
        RevaluationLayer = self.env["stock.valuation.layer.revaluation"]
        for line in self:
            product = line.product_id.sudo().with_company(line.adjustment_id.company_id)
            qty_svl = product.quantity_svl
            if qty_svl <= 0:
                continue
            diff = (qty_svl * line.real_cost) - product.value_svl
            valuation_layer_model = self.env["stock.valuation.layer"].sudo()
            previous_layer_ids = valuation_layer_model.search(
                [
                    ("product_id", "=", product.id),
                    "|",
                    ("company_id", "=", self.env.companies.ids),
                    ("company_id", "=", False),
                ]
            ).ids
            layer = RevaluationLayer.create(
                {
                    "product_id": product.id,
                    "added_value": diff,
                    "reason": line.adjustment_id.reason_id.name,
                    "account_id": line.adjustment_id.account_id.id,
                    "date": line.adjustment_id.date,
                    "company_id": line.adjustment_id.company_id.id,
                }
            )
            res = layer.action_validate_revaluation()

            new_layer = valuation_layer_model.search(
                [
                    ("product_id", "=", product.id),
                    "|",
                    ("company_id", "=", self.env.companies.ids),
                    ("company_id", "=", False),
                    ("id", "not in", previous_layer_ids),
                ],
                order="id desc",
                limit=1,
            )
            if new_layer.account_move_id:
                line.account_move_ids = [Command.link(new_layer.account_move_id.id)]
                line.adjustment_id.account_move_ids = [
                    Command.link(new_layer.account_move_id.id)
                ]

            if res:
                line.env.user.notify_success(_("All Cost Updated"))

    @api.model
    def _get_adjustment_src_location(self, picking_type, account_id=False):
        """Location used as source of the adjustment receipt.

        If a counterpart account is set on the adjustment, use (or create) a
        child location of the picking type default source with that account as
        'valuation_out_account_id', so the journal entry credits it instead of
        the category's stock input account.
        """
        default_src = picking_type.default_location_src_id
        if not default_src:
            default_src = self.env.ref("stock.stock_location_suppliers", False)
        if not default_src or not account_id:
            return default_src
        loc = self.env["stock.location"].search(
            [
                ("name", "=", "Ajustes de Inventario"),
                ("location_id", "=", default_src.id),
            ],
            limit=1,
        )
        if not loc:
            loc = self.env["stock.location"].create(
                {
                    "name": "Ajustes de Inventario",
                    "location_id": default_src.id,
                }
            )
        if loc.valuation_out_account_id != account_id:
            loc.valuation_out_account_id = account_id.id
        return loc

    def _process_qty_cost(self):
        for line in self:
            adj = line.adjustment_id
            dest_loc = adj.origin_location_id
            qty_to_receive = line.new_qty
            if qty_to_receive <= 0:
                continue

            # Warehouse whose stock hierarchy contains the destination location
            warehouse = self.env["stock.warehouse"].search(
                [
                    ("lot_stock_id", "parent_of", dest_loc.id),
                    ("company_id", "=", adj.company_id.id),
                ],
                limit=1,
            )
            in_type = (
                warehouse.in_type_id
                if warehouse
                else self.env["stock.picking.type"].search(
                    [
                        ("code", "=", "incoming"),
                        ("default_location_src_id", "!=", False),
                        ("warehouse_id.pf_branch_id", "=", dest_loc.pf_branch_id.id),
                    ],
                    limit=1,
                )
            )
            if not in_type or not (
                in_type.default_location_src_id
                or self.env.ref("stock.stock_location_suppliers", False)
            ):
                raise UserError(
                    _(
                        "No incoming operation type with a source location could be "
                        "found for location '%(dest_loc)s' "
                        "(branch '%(dest_loc_branch)s'). Check the warehouse "
                        "configuration."
                    )
                    % {
                        "dest_loc": dest_loc.complete_name,
                        "dest_loc_branch": dest_loc.pf_branch_id.name or "",
                    }
                )

            src_loc = self._get_adjustment_src_location(in_type, adj.account_id)
            if not src_loc:
                raise UserError(
                    _(
                        "The operation type '%s' has no default source location "
                        "and the 'Vendors' location is missing."
                    )
                    % in_type.name
                )

            picking = self.env["stock.picking"].create(
                {
                    "picking_type_id": in_type.id,
                    "location_id": src_loc.id,
                    "location_dest_id": dest_loc.id,
                    "company_id": adj.company_id.id,
                    "pf_branch_id": dest_loc.pf_branch_id.id,
                    "move_ids_without_package": [
                        (
                            0,
                            0,
                            {
                                "name": line.product_id.display_name,
                                "product_id": line.product_id.id,
                                "product_uom": line.product_uom_id.id,
                                "product_uom_qty": qty_to_receive,
                                "location_id": src_loc.id,
                                "location_dest_id": dest_loc.id,
                                "price_unit": line.real_cost,
                            },
                        )
                    ],
                }
            )

            line.adjustment_id.picking_ids = [Command.link(picking.id)]

            if hasattr(picking, "action_reserve_picking"):
                picking.action_reserve_picking()
                picking.action_confirm_picking()

            picking.action_confirm()
            picking.move_ids_without_package._set_quantity_done(qty_to_receive)
            picking.button_validate()

            account_moves = picking.move_ids_without_package.mapped(
                "account_move_ids"
            ) | picking.move_ids_without_package.mapped(
                "stock_valuation_layer_ids.account_move_id"
            )
            line.stock_move_ids = [
                Command.link(move.id) for move in picking.move_ids_without_package
            ]
            if account_moves:
                line.account_move_ids = [
                    Command.link(move.id) for move in account_moves
                ]
                adj.account_move_ids = [Command.link(move.id) for move in account_moves]
